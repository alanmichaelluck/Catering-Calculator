import streamlit as st
import openpyxl
import pandas as pd
import re
from datetime import datetime
from io import BytesIO

st.set_page_config(page_title="Catering + Prep Report Integrator", layout="wide")
st.title("📋 Unified Prep Report & Multi-Day Catering Calculator")

# ==============================================================================
# EXACT KITCHEN CROSSWALK (Based on your master mapping sheet)
# ==============================================================================
# Maps each Catering Calculator item to its Prep Report sub-recipes & scaling factors.
# Items marked "Ignore" are intentionally omitted here so they keep default calculations.
CATERING_CROSSWALK = {
    # Mac: .3 of a 1/6 pan cheese sauce and 12% of a 4" hotel pan GF pasta (per 1x catering unit)
    "Mac": [
        {"prep_recipe": "Cheese Sauce (Flower)", "source_col": "Catering_Qty", "multiplier": 0.3},
        {"prep_recipe": "Gluten Free Pasta (Flower)", "source_col": "Catering_Qty", "multiplier": 0.12},
    ],
    # Quinoa Corn: 6 portions
    "Quinoa Corn": [
        {"prep_recipe": "Grains - Red Quinoa Prepped (Flower)", "source_col": "Catering_Qty", "multiplier": 6.0},
    ],
    # Sweets: 700g roasted sweet potato (~1.54 Lbs or 1.5Qt equivalent)
    "Sweets": [
        {"prep_recipe": "Roasted Sweet Potatoes (Flower)", "source_col": "Catering_Qty", "multiplier": 1.54},
    ],
    # Smash: 6 portions of Smashed Potatoes
    "Smash": [
        {"prep_recipe": "Smashed Potatoes Prep (Flower)", "source_col": "Catering_Qty", "multiplier": 6.0},
    ],
    # Broccoli: 6.4 portions of roasted broccoli
    "Broccoli": [
        {"prep_recipe": "Roasted Broccoli Florets - Prep (Flower)", "source_col": "Catering_Qty", "multiplier": 6.4},
    ],
    # Fruit: 5.7 portions
    "Fruit": [
        {"prep_recipe": "Fruit Salad Mix (Flower)", "source_col": "Catering_Qty", "multiplier": 5.7},
    ],
    # Brown Rice: 20% of a 1x batch (1x batch = 10 Qt, so 2 Qt per 1x catering)
    "Brown Rice": [
        {"prep_recipe": "Brown Rice Prep (Flower)", "source_col": "Catering_Qty", "multiplier": 2.0},
    ],
    # Jasmine Rice: 20% of a 1x batch (1x batch = 16 portions, so 3.2 portions per 1x catering)
    "Jasmine Rice": [
        {"prep_recipe": "Jasmine Rice Prep (Flower)", "source_col": "Catering_Qty", "multiplier": 3.2},
    ],
    # Cauli Risotto: 1qt diced cauli, 1qt cauli cream
    "Cauli Risotto": [
        {"prep_recipe": "Cauliflower Dice & Blanched (Flower)", "source_col": "Catering_Qty", "multiplier": 1.0},
        {"prep_recipe": "Cauliflower Cream (Flower)", "source_col": "Catering_Qty", "multiplier": 1.0},
    ],
    # Olive Oil Veg: 7.5oz olive oil roasted cauli (~0.23 Qt)
    "Olive Oil Veg": [
        {"prep_recipe": "Olive Oil Roasted Cauliflower (Flower)", "source_col": "Catering_Qty", "multiplier": 0.23},
    ],
    # Brussels Sprouts: 17.5oz roasted brussel sprouts (~1.1 Lb)
    "Brussels Sprouts": [
        {"prep_recipe": "Roasted Brussels Sprouts (Flower)", "source_col": "Catering_Qty", "multiplier": 1.1},
    ],
    # Flo Cho Chow: 8oz flo cho chow sauce (= 0.25 Qt)
    "Flo Cho Chow": [
        {"prep_recipe": "Flo Cho Chow Sauce (Flower)", "source_col": "Catering_Qty", "multiplier": 0.25},
    ],
    # Forbidden: 10 fl oz sesame hoisin sauce (= 0.3125 Qt)
    "Forbidden": [
        {"prep_recipe": "Sesame Hoisin Sauce (Flower)", "source_col": "Catering_Qty", "multiplier": 0.3125},
    ],
    # Chicken Yakisoba: 8oz garlic noodle sauce (= 0.25 Qt)
    "Chicken Yakisoba": [
        {"prep_recipe": "Garlic Noodle Sauce (Flower)", "source_col": "Catering_Qty", "multiplier": 0.25},
    ],
    # Salads: 4 portions of respective salad mix + dressing
    "Skinny": [
        {"prep_recipe": "Salad Mix - Skinny Cobb (Flower)", "source_col": "Catering_Qty", "multiplier": 4.0},
        {"prep_recipe": "Smoked Blue Cheese Vinaigrette (Flower)", "source_col": "Catering_Qty", "multiplier": 0.25}, # 8 oz / 32 oz per Qt
    ],
    "Kale": [
        {"prep_recipe": "Salad Mix - Ginger Miso Crunch (Flower)", "source_col": "Catering_Qty", "multiplier": 4.0},
        {"prep_recipe": "Ginger Miso Vinaigrette (Flower)", "source_col": "Catering_Qty", "multiplier": 0.25},
    ],
    "Chopped": [
        {"prep_recipe": "Salad Mix - Chopped Veg (Flower)", "source_col": "Catering_Qty", "multiplier": 4.0},
        {"prep_recipe": "Oregano Dressing (Flower)", "source_col": "Catering_Qty", "multiplier": 0.25},
    ],
    "Avo Caesar": [
        {"prep_recipe": "Avocado Chicken Caesar Salad (Flower)", "source_col": "Catering_Qty", "multiplier": 4.0},
        {"prep_recipe": "Avocado Caesar Dressing (Flower)", "source_col": "Catering_Qty", "multiplier": 0.25},
    ],
}

# Catering Calculator item name row mapping on Excel sheet
CATERING_NAME_MAP = {
    "Mac": 4, "Quinoa Corn": 5, "Sweets": 6, "Smash": 7, "Broccoli": 8, "Fruit": 9,
    "Brown Rice": 10, "Jasmine Rice": 11, "Butternut Gratin": 12, "Kale Salad": 13,
    "Cauli Risotto": 14, "Olive Oil Veg": 15, "Brussels Sprouts": 16,
    "Flo Cho Chow": 18, "Forbidden": 19, "Chicken Enchiladas": 20, "Miso Glazed Salmon": 21,
    "Chicken Yakisoba": 22, "Mother": 23,
    "Skinny": 25, "Kale": 27, "Chopped": 29, "Avo Caesar": 31, "Butternut & Pear": 33,
    "Flying Avo": 43, "The 1924": 44, "Tofu \"Banh Mi\"": 45, "Rebel": 46,
    "Chicken": 48, "Steak": 49, "Salmon": 50, "Tofu": 51, "Shrimp": 52,
    "Cookies": 54, "Brownies": 55, "1/2 & 1/2": 56
}

# ==============================================================================
# HELPER FUNCTIONS
# ==============================================================================
def extract_catering_demands(cat_file_bytes):
    wb = openpyxl.load_workbook(BytesIO(cat_file_bytes), data_only=True)
    ws = wb['Catering Calculator']
    raw_date = ws.cell(1, 8).value
    cat_date = raw_date.date() if isinstance(raw_date, datetime) else pd.to_datetime(str(raw_date)).date()

    demands = {}
    for cat_item, rules in CATERING_CROSSWALK.items():
        row_idx = CATERING_NAME_MAP.get(cat_item)
        if row_idx:
            catering_qty = ws.cell(row_idx, 4).value or 0.0 # Column D (Catering Qty multiplier e.g. 8)
            if isinstance(catering_qty, (int, float)) and catering_qty > 0:
                for rule in rules:
                    prep_rec = rule["prep_recipe"]
                    multiplier = rule["multiplier"]
                    added_qty = float(catering_qty) * multiplier
                    demands[prep_rec] = demands.get(prep_rec, 0.0) + added_qty
    return cat_date, demands

def parse_shelf_life_days(sl_str):
    s = str(sl_str).strip().lower()
    if s.endswith('h'): return 1
    if s.endswith('d'):
        days = int(re.findall(r'\d+', s)[0])
        return min(max(days, 1), 3)
    return 1

def parse_batch_size(bs_str, iu_str):
    bs = str(bs_str).strip()
    iu = str(iu_str).strip()
    if bs == '1 X = 4" Hotel Pan': return 1.0
    if bs == '1 X = 0.75 Qt' and iu == 'Cup': return 3.0
    if '1 X =' in bs:
        after = bs.split('1 X =')[1].strip()
        m = re.match(r'^(\d+(?:\.\d+)?)', after)
        if m: return float(m.group(1))
    return 1.0

# ==============================================================================
# STREAMLIT UI
# ==============================================================================
col1, col2 = st.columns(2)
with col1:
    prep_file = st.file_uploader("1. Upload Next-Day Prep Report (.xlsx)", type=["xlsx"])
with col2:
    cat_files = st.file_uploader("2. Upload ALL 3 Catering Calculator Files (.xlsx)", type=["xlsx"], accept_multiple_files=True)

if prep_file and cat_files:
    wb_prep = openpyxl.load_workbook(prep_file)
    ws_prep = wb_prep['Prep Report']

    raw_prep_date = str(ws_prep.cell(4, 9).value)
    date_match = re.search(r'(\d{1,2}/\d{1,2}/\d{4})', raw_prep_date)
    prep_date = datetime.strptime(date_match.group(1), '%m/%d/%Y').date()
    st.info(f"**Prep Report Target Date (Day 1):** {prep_date.strftime('%A, %m/%d/%Y')}")

    day_demands = {1: {}, 2: {}, 3: {}}
    file_summary = []
    
    for f in cat_files:
        c_date, c_demands = extract_catering_demands(f.getvalue())
        day_offset = (c_date - prep_date).days + 1
        file_summary.append(f"📁 **{f.name}** -> Date: **{c_date.strftime('%m/%d/%Y')}** (Assigned to **Day {day_offset}**)")
        if day_offset in [1, 2, 3]:
            for rec, qty in c_demands.items():
                day_demands[day_offset][rec] = day_demands[day_offset].get(rec, 0.0) + qty

    st.success("\n".join(file_summary))

    preview_rows = []
    for r in range(6, ws_prep.max_row + 1):
        rec_name = ws_prep.cell(r, 2).value
        sl_str = ws_prep.cell(r, 4).value
        bs_str = ws_prep.cell(r, 7).value
        iu_str = ws_prep.cell(r, 8).value
        orig_dc = ws_prep.cell(r, 11).value
        orig_slm = ws_prep.cell(r, 12).value

        if sl_str is not None and rec_name is not None:
            orig_dc_val = float(orig_dc or 0.0)
            orig_slm_val = float(orig_slm or 0.0)
            sl_days = parse_shelf_life_days(sl_str)
            batch_yield = parse_batch_size(bs_str, iu_str)

            d1_cat = day_demands[1].get(rec_name, 0.0)
            window_cat = sum(day_demands[d].get(rec_name, 0.0) for d in range(1, sl_days + 1))
            window_cat_batches = window_cat / batch_yield if batch_yield > 0 else 0.0

            new_dc_val = round((orig_dc_val + d1_cat) * 4) / 4
            new_slm_val = round((orig_slm_val + window_cat_batches) * 4) / 4

            # ONLY update rows that have catering demand attached; leave everything else untouched!
            if d1_cat > 0 or window_cat > 0:
                ws_prep.cell(r, 11).value = new_dc_val
                ws_prep.cell(r, 12).value = new_slm_val

                preview_rows.append({
                    "Recipe Name": rec_name,
                    "Shelf Life": sl_str,
                    "Inv. Unit": iu_str,
                    "Orig Daily Cons": orig_dc_val,
                    "+ Day 1 Catering": round(d1_cat, 2),
                    "NEW Daily Cons": new_dc_val,
                    "Orig Shelf Life Max (X)": orig_slm_val,
                    f"+ Catering Batches ({sl_days}d)": round(window_cat_batches, 2),
                    "NEW Shelf Life Max (X)": new_slm_val,
                })

    st.subheader("🔍 Recipes Updated by Catering (All Other Recipes Kept As-Is)")
    st.dataframe(pd.DataFrame(preview_rows), use_container_width=True)

    output = BytesIO()
    wb_prep.save(output)
    st.download_button(
        label="📥 Download Updated Walk-In Prep Report (.xlsx)",
        data=output.getvalue(),
        file_name=f"Prep_Report_With_Catering_{prep_date.strftime('%Y%m%d')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
